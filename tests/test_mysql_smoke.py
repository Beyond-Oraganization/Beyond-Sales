"""Test dymny warstwy danych na PRAWDZIWYM MySQL (to, czego sqlite nie pokryje:
tryb ścisły, utf8mb4, słowo zastrzeżone `row_number`, FOR UPDATE, strumieniowanie).

Uruchamiany wyłącznie, gdy ustawiono PERNO_TEST_DB_NAME albo PERNO_TEST_DB_URL, np.:
    PERNO_TEST_DB_NAME=perno_maile_test python3 -m unittest tests.test_mysql_smoke -v
    PERNO_TEST_DB_URL=mysql+pymysql://root@127.0.0.1:33307/perno_maile_test .venv/bin/python -m pytest tests/test_mysql_smoke.py
Nazwa bazy MUSI kończyć się na `_test` — test usuwa i odtwarza w niej tabele.
Z PERNO_TEST_DB_NAME dane logowania bierze z config.env (DB_HOST/DB_PORT/DB_USER/DB_PASSWORD); PERNO_TEST_DB_URL
podaje cały adres jawnie (np. do tymczasowego serwera z `scripts/scratch_mysql.sh`) i config.env nie jest czytany.

Druga klasa (`MysqlScraperSmokeTests`) sprawdza tabele i zapytania skrapera na prawdziwym MySQL: kolumny ASCII,
upserty `ON DUPLICATE KEY UPDATE`, BLOB, `FOR UPDATE`, wyłączność dzierżawy między wątkami, `UNIQUE` z NULL.
"""

from __future__ import annotations

import io
import os
import threading
import unittest
from datetime import datetime, timedelta

import openpyxl
import sqlalchemy as sa
from sqlalchemy import exc as sa_exc

from perno import db as dbmod
from perno import settings as cfg
from perno.domain import statuses as st
from perno.domain.rowfilter import NO_COUNTRY, RowFilter
from perno.scraper import schema as scraper_schema
from perno.scraper import statuses as sst
from perno.scraper.contacts import Contact
from perno.scraper.store import LeaseLost, LinkRecord, PageRecord, RunActive, Store
from perno.scraper.urls import url_hash
from perno.services import bulk
from perno.services import export as exporter
from tests.fixtures import put as put_rows
from tests.fixtures import sample_row

TEST_URL = os.environ.get("PERNO_TEST_DB_URL", "")
TEST_DB = (sa.engine.make_url(TEST_URL).database if TEST_URL else os.environ.get("PERNO_TEST_DB_NAME", "")) or ""
SKIP_REASON = "ustaw PERNO_TEST_DB_NAME=perno_maile_test (albo PERNO_TEST_DB_URL), żeby uruchomić test na prawdziwym MySQL"


def resolve_test_url() -> sa.engine.URL:
    """Adres bazy testowej (jawny PERNO_TEST_DB_URL albo config.env + PERNO_TEST_DB_NAME); baza musi się kończyć na `_test`."""
    if not TEST_DB.endswith("_test"):
        raise RuntimeError("Nazwa bazy testowej musi kończyć się na _test (test kasuje w niej tabele).")
    if TEST_URL:
        return sa.engine.make_url(TEST_URL)
    config = cfg.load_config(cfg.DEFAULT_CONFIG)
    return dbmod.build_url({**config, "DB_NAME": TEST_DB}, env=cfg.Environment())


@unittest.skipUnless(TEST_DB, SKIP_REASON)
class MysqlSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.database = dbmod.Database(dbmod.make_engine(resolve_test_url()))
        cls.info = cls.database.check_connection()
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)
        cls.database.ensure_schema()

    @classmethod
    def tearDownClass(cls):
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)

    def setUp(self):
        with self.database.transaction() as conn:
            conn.execute(dbmod.maile.delete())
            conn.execute(dbmod.companies.delete())

    def put(self, *rows):
        with self.database.transaction() as conn:
            self.database.insert_rows(conn, list(rows))

    def test_server_is_strict_and_is_mysql(self):
        self.assertEqual(self.info["backend"], "mysql")
        self.assertIn("STRICT", self.info["sql_mode"])

    def test_search_collation_is_checked_and_a_binary_collation_is_reported(self):
        self.assertEqual(self.database.check_connection()["search_warning"], "")
        with self.database.engine.begin() as conn:
            conn.execute(sa.text("ALTER TABLE maile MODIFY company VARCHAR(255) COLLATE utf8mb4_bin"))
        warning = self.database.check_connection()["search_warning"]
        self.assertIn("company (utf8mb4_bin)", warning)
        self.assertIn("CONVERT TO CHARACTER SET", warning)
        with self.database.engine.begin() as conn:
            conn.execute(sa.text("ALTER TABLE maile MODIFY company VARCHAR(255) COLLATE utf8mb4_0900_ai_ci"))
        self.assertEqual(self.database.check_connection()["search_warning"], "")

    def test_utf8mb4_and_long_bodies_roundtrip(self):
        body = "Zażółć gęślą jaźń 🔗 — " * 5000
        self.put(sample_row(1, company="Łódź 🚀", body_html=body, body_text=body))
        row = self.database.get("eventy-001")
        self.assertEqual((row["company"], row["body_html"]), ("Łódź 🚀", body))
        self.assertGreater(len(row["body_html"]), 70_000)

    def test_datetime_has_second_precision_and_reserved_word_column_works(self):
        self.put(sample_row(7, status="wyslano", sent_at="2026-09-29T14:23:11"))
        row = self.database.get("eventy-007")
        self.assertEqual((row["sent_at"], row["row_number"]), ("2026-09-29T14:23:11", "7"))
        self.assertEqual(self.database.fetch_all()[0]["id"], "eventy-007")

    def test_claim_has_exactly_one_winner_among_threads(self):
        self.put(sample_row(1))
        winners, barrier = [], threading.Barrier(8)

        def worker():
            barrier.wait()
            if self.database.claim_for_sending("eventy-001"):
                winners.append(1)

        threads = [threading.Thread(target=worker) for _ in range(8)]
        [t.start() for t in threads]
        [t.join() for t in threads]
        self.assertEqual(len(winners), 1)

    def test_edit_upsert_and_bookkeeping(self):
        self.put(sample_row(1, status="pominieto_brak_email", email=""), sample_row(2, status="wyslano", sent_at="2026-09-29T10:00:00"))
        self.assertEqual(self.database.apply_edit("eventy-001", {"email": "a@b.pl"})["status"], st.STATUS_PENDING)
        stats = self.database.upsert_generated([sample_row(1, subject="Nowy"), sample_row(2, subject="Nie ruszaj"), sample_row(3)])
        self.assertEqual(stats, {"inserted": 1, "updated": 1, "skipped": 1})
        self.assertEqual(self.database.get("eventy-002")["subject"], "Temat 2")
        self.assertEqual(self.database.sent_today_count(datetime(2026, 9, 29, 12, 0, 0)), 1)

    def test_streaming_export_against_mysql(self):
        self.put(*[sample_row(n, company=f"Firma {n}") for n in range(1, 301)])
        for fmt in ("csv", "xlsx"):
            result = exporter.build_export(self.database, exporter.ExportSpec(fmt=fmt, bodies=True))
            with result.stream as f:
                data = f.read()
            if fmt == "xlsx":
                self.assertEqual(openpyxl.load_workbook(io.BytesIO(data))["Maile"].max_row, 301)
            else:
                self.assertEqual(data.decode("utf-8-sig").count("Firma "), 300)

    def test_bulk_update_and_undo_against_mysql(self):
        """Zbiorcza zmiana: SELECT … FOR UPDATE w kolejności id, UPDATE w jednej transakcji i cofanie."""
        self.put(sample_row(1, status="wyslano", sent_at="2026-09-29T14:00:00", country="PL"),
                 sample_row(2), sample_row(3, status="wysylanie"))
        log = bulk.UndoLog()
        request = bulk.parse_request({"ids": ["eventy-003", "eventy-001", "eventy-002"],
                                      "changes": {"status": "odpowiedzial", "city": "Łódź"}})
        result = bulk.apply(self.database, request, log)
        self.assertEqual((result["changed"], result["skipped_count"]), (2, 1))
        row = self.database.get("eventy-001")
        self.assertEqual((row["status"], row["sent_at"], row["city"]), ("odpowiedzial", "2026-09-29T14:00:00", "Łódź"))
        self.assertEqual(self.database.get("eventy-003")["status"], "wysylanie")
        undone = bulk.undo(self.database, log)
        self.assertEqual((undone["restored"], undone["conflict_count"]), (2, 0))
        self.assertEqual((self.database.get("eventy-001")["status"], self.database.get("eventy-001")["city"]), ("wyslano", ""))
        with self.assertRaises(bulk.AckRequired):
            bulk.apply(self.database, bulk.parse_request({"ids": ["eventy-001"], "changes": {"status": "do_wyslania"}}), log)
        self.assertEqual(self.database.get("eventy-001")["status"], "wyslano")

    def test_overlapping_bulk_updates_serialize_without_deadlock(self):
        self.put(*[sample_row(n) for n in range(1, 61)])
        ids = [f"eventy-{n:03d}" for n in range(1, 61)]
        errors, log = [], bulk.UndoLog()

        def run(status, order):
            try:
                request = bulk.parse_request({"ids": order, "changes": {"status": status}})
                bulk.apply(self.database, request, log)
            except Exception as exc:  # noqa: BLE001
                errors.append(exc)

        threads = [threading.Thread(target=run, args=("odpowiedzial", ids)),
                   threading.Thread(target=run, args=("zainteresowany", list(reversed(ids))))]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=60)
        self.assertEqual(errors, [])
        # blokady zakładane zawsze w kolejności id, więc jedna operacja kończy się w całości przed drugą
        self.assertEqual(len({row["status"] for row in self.database.fetch_all()}), 1)

    def test_filter_dimensions_light_list_and_ordered_ids_against_mysql(self):
        self.put(sample_row(1, country="PL", email="a@x.pl"), sample_row(2, country="GB", email="   ", body_html="", body_text=""),
                 sample_row(3, country="", email=""), sample_row(4, status="wyslano", sent_at="2026-09-29T23:59:59", country="DE"))
        ids = lambda **kw: [r["id"] for r in self.database.iter_rows(RowFilter(**kw), ("id",))]
        self.assertEqual(sorted(ids(countries=("PL", NO_COUNTRY))), ["eventy-001", "eventy-003"])
        self.assertEqual(sorted(ids(has_email=False)), ["eventy-002", "eventy-003"])  # spacje to brak adresu
        self.assertEqual(ids(has_body=False), ["eventy-002"])
        from datetime import date
        self.assertEqual(ids(sent_from=date(2026, 9, 29), sent_to=date(2026, 9, 29)), ["eventy-004"])
        self.assertEqual(ids(ids=("eventy-004", "nie-ma", "eventy-001")), ["eventy-004", "eventy-001"])
        light = {r["id"]: r for r in self.database.fetch_all(light=True)}
        self.assertTrue(light["eventy-001"]["has_body"])
        self.assertFalse(light["eventy-002"]["has_body"])
        self.assertNotIn("body_html", light["eventy-001"])



@unittest.skipUnless(TEST_DB, SKIP_REASON)
class MysqlScraperSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.database = dbmod.Database(dbmod.make_engine(resolve_test_url()))
        scraper_schema.skraper_metadata.drop_all(cls.database.engine)
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)
        cls.database.ensure_schema()
        cls.store = Store(cls.database)
        cls.store.ensure_schema()

    @classmethod
    def tearDownClass(cls):
        scraper_schema.skraper_metadata.drop_all(cls.database.engine)
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)

    def setUp(self):
        with self.database.transaction() as conn:
            for table in reversed(scraper_schema.skraper_metadata.sorted_tables):
                conn.execute(table.delete())
            conn.execute(dbmod.maile.delete())
            conn.execute(dbmod.companies.delete())

    def firm(self, n=1):
        return self.store.add_manual(f"https://firma{n}.example.com/", f"Firma {n}")["id"]

    def lease(self, firma_id):
        run_id = self.store.start_run({})
        self.store.enqueue([firma_id])
        self.assertEqual(self.store.claim_next_firma(run_id)["id"], firma_id)
        return run_id

    def page(self, path="/", **kw):
        url = "https://firma1.example.com" + path
        values = dict(url=url, url_hash=url_hash(url), depth=0, priority=0, source="start", status="pobrano",
                      http_status=200, text="treść", title="Tytuł", bytes=1000)
        values.update(kw)
        return PageRecord(**values)

    def test_url_columns_are_ascii_and_reject_other_characters_in_strict_mode(self):
        with self.database.engine.connect() as conn:
            charset = conn.execute(sa.text(
                "SELECT CHARACTER_SET_NAME FROM information_schema.COLUMNS "
                "WHERE TABLE_SCHEMA = :db AND TABLE_NAME = 'strony' AND COLUMN_NAME = 'url'"), {"db": TEST_DB}).scalar()
        self.assertEqual(charset, "ascii")
        firma_id = self.firm()
        with self.assertRaises(sa_exc.DBAPIError):
            with self.database.transaction() as conn:
                conn.execute(sa.insert(scraper_schema.strony).values(
                    firma_id=firma_id, url="https://firma1.example.com/zażółć", url_hash="a" * 40, glebokosc=0,
                    priorytet=0, zrodlo="link", status="oczekuje"))

    def test_noop_upsert_ignores_duplicates_without_masking_other_errors(self):
        firma_id = self.firm()
        run_id = self.lease(firma_id)
        link = LinkRecord("https://firma1.example.com/a", url_hash("https://firma1.example.com/a"), 1, 10)
        self.store.add_discovered(firma_id, run_id, [link])
        self.store.add_discovered(firma_id, run_id, [link, link])
        self.assertEqual(self.store.pending_count(firma_id), 1)

    def test_save_result_roundtrip_with_emoji_nul_blob_and_contact_counter(self):
        firma_id = self.firm()
        run_id = self.lease(firma_id)
        blob = os.urandom(2 * 1024 * 1024)
        contact = Contact("email", "biuro@firma1.example.com", "mailto", True, "ogolny")
        self.store.save_result(firma_id, run_id, self.page(
            "/", text="Zażółć 🚀\x00 tekst", html_gz=blob, contacts=[contact],
            links=[LinkRecord("https://firma1.example.com/x", url_hash("https://firma1.example.com/x"), 1, 5)]))
        self.store.save_result(firma_id, run_id, self.page("/dwa", contacts=[contact]))
        with self.database.engine.connect() as conn:
            content = conn.execute(sa.select(scraper_schema.strony_tresc.c.tekst, scraper_schema.strony_tresc.c.html_gz)
                                   .order_by(scraper_schema.strony_tresc.c.strona_id)).first()
            counter = conn.execute(sa.select(scraper_schema.kontakty.c.liczba_stron)).scalar()
        self.assertEqual(content.tekst, "Zażółć 🚀 tekst")
        self.assertEqual(bytes(content.html_gz), blob)
        self.assertEqual(counter, 2)
        firm = self.store.get_firma(firma_id)
        self.assertEqual((firm["stron_pobranych"], firm["bajty_pobranych"]), (2, 2000))

    def test_lease_is_exclusive_and_fenced(self):
        ids = [self.firm(n) for n in range(1, 31)]
        run_id = self.store.start_run({})
        self.store.enqueue()
        got, lock, barrier = [], threading.Lock(), threading.Barrier(8)

        def worker():
            barrier.wait()
            while True:
                firm = self.store.claim_next_firma(run_id)
                if firm is None:
                    return
                with lock:
                    got.append(firm["id"])

        threads = [threading.Thread(target=worker) for _ in range(8)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(60)
        self.assertEqual(sorted(got), sorted(ids))
        self.assertTrue(self.store.release_firma(ids[0], run_id, sst.FIRMA_ZAKONCZONA))
        with self.assertRaises(LeaseLost):
            self.store.save_result(ids[0], run_id, self.page("/", url="https://firma1.example.com/"))

    def test_only_one_active_run_and_null_unique_semantics(self):
        first = self.store.start_run({})
        with self.assertRaises(RunActive):
            self.store.start_run({})
        self.store.finish_run(first, sst.RUN_ZAKONCZONY)
        second = self.store.start_run({})
        self.store.finish_run(second, sst.RUN_ZAKONCZONY)
        third = self.store.start_run({})  # wiele wierszy z aktywny = NULL jest dozwolone
        self.assertTrue(first < second < third)

    def test_heartbeat_and_stop_do_not_depend_on_rowcount(self):
        run_id = self.store.start_run({})
        self.assertFalse(self.store.heartbeat(run_id, {"a": 1}))
        self.assertFalse(self.store.heartbeat(run_id, {"a": 1}))  # ten sam stan w tej samej sekundzie
        self.assertTrue(self.store.request_stop())
        self.assertTrue(self.store.request_stop())
        self.assertTrue(self.store.heartbeat(run_id))

    def test_seed_delete_and_purge_with_subqueries(self):
        with self.database.transaction() as conn:
            self.database.insert_rows(conn, [sample_row(1), sample_row(2)])
        self.assertEqual(self.store.seed_from_maile()["dodano"], 2)
        firma_id = self.store.find_firma("firma1.example.com")["id"]
        run_id = self.lease(firma_id)
        self.store.save_result(firma_id, run_id, self.page("/", html_gz=b"x" * 100))
        self.store.release_firma(firma_id, run_id, sst.FIRMA_ZAKONCZONA)
        self.assertEqual(self.store.purge_html(-1), 1)  # UPDATE … WHERE strona_id IN (SELECT …)
        self.store.delete_firma_data(firma_id)  # DELETE … WHERE strona_id IN (SELECT …)
        with self.database.engine.connect() as conn:
            self.assertEqual(conn.execute(sa.select(sa.func.count()).select_from(scraper_schema.strony)).scalar(), 0)
        self.assertEqual(self.store.get_firma(firma_id)["skan_status"], sst.FIRMA_NIESKANOWANA)

    def test_robots_text_and_long_urls_fit(self):
        firma_id = self.firm()
        run_id = self.lease(firma_id)
        self.store.save_robots(firma_id, run_id, "ok", "User-agent: *\n" + "# zażółć\n" * 20000, 1.5)
        url = "https://firma1.example.com/" + "a" * 1900
        self.store.save_result(firma_id, run_id, self.page("/x", url=url, url_hash=url_hash(url)))
        self.assertEqual(self.store.get_firma(firma_id)["crawl_delay"], 1.5)
        self.assertGreater(len(self.store.get_firma(firma_id)["robots_txt"]), 150_000)


MYSQL_OLD_TABLE = """
CREATE TABLE maile (
    id VARCHAR(32) NOT NULL PRIMARY KEY, segment VARCHAR(32) NOT NULL, `row_number` INT, company VARCHAR(255),
    contact_person VARCHAR(255), email VARCHAR(320), website VARCHAR(1000), tag VARCHAR(255), mode VARCHAR(16),
    subject VARCHAR(500), body_html MEDIUMTEXT, body_text MEDIUMTEXT, status VARCHAR(32) NOT NULL, generated_at DATETIME,
    sent_at DATETIME, attachments TEXT, inline_images VARCHAR(255), error TEXT,
    retry_count INT NOT NULL DEFAULT 0, updated_at DATETIME, country VARCHAR(2), city VARCHAR(120),
    address VARCHAR(500), description TEXT
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci
"""


@unittest.skipUnless(TEST_DB, SKIP_REASON)
class MysqlAdoptionTests(unittest.TestCase):
    """Uaktualnienie i przejęcie na prawdziwym MySQL: DDL, indeks unikalny, `GET_LOCK`, `INSERT … SELECT`, porównanie
    z NULL (`<=>`) i zachowanie przy dwóch jednoczesnych przejęciach."""

    def setUp(self):
        self.database = dbmod.Database(dbmod.make_engine(resolve_test_url()))
        dbmod.metadata.drop_all(self.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)
        with self.database.engine.begin() as conn:
            conn.execute(sa.text(MYSQL_OLD_TABLE))
            for n in (1, 2, 3):
                conn.execute(sa.text(
                    "INSERT INTO maile (id, segment, `row_number`, company, email, status, retry_count, country) "
                    "VALUES (:id, 'eventy', :n, :company, :email, 'wyslano', 0, 'PL')"),
                    {"id": f"eventy-{n:03d}", "n": n, "company": f"Firma {n}", "email": f"f{n}@example.com"})
        self.addCleanup(dbmod.metadata.drop_all, self.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)

    def test_upgrade_adopt_and_diagnose(self):
        added = self.database.upgrade_schema()
        self.assertIn("podmioty", added)
        self.assertEqual(added[-1], "uq_maile_company_email")
        self.assertEqual(self.database.upgrade_schema(), [])
        report = self.database.adopt()
        self.assertEqual((report.companies_added, report.rows_linked), (3, 3))
        self.assertTrue(self.database.diagnose().clean, self.database.diagnose())
        self.assertEqual(self.database.adopt().rows_linked, 0)
        with self.database.engine.connect() as conn:
            self.assertEqual(conn.execute(sa.text("SELECT IS_FREE_LOCK('perno_adopt')")).scalar(), 1)  # blokada zwolniona

    def test_the_unique_index_refuses_a_duplicate_address_in_one_company(self):
        self.database.upgrade_schema()
        self.database.adopt()
        with self.assertRaises(sa_exc.IntegrityError), self.database.engine.begin() as conn:
            conn.execute(sa.text("INSERT INTO maile (id, segment, email, status, retry_count, company_id, position) "
                                 "VALUES ('eventy-001~2', 'eventy', 'F1@EXAMPLE.COM', 'do_wyslania', 0, 'eventy-001', 2)"))

    def test_mirror_drift_is_detected_with_null_safe_comparison(self):
        self.database.upgrade_schema()
        self.database.adopt()
        with self.database.engine.begin() as conn:
            conn.execute(sa.text("UPDATE maile SET city = 'Łódź' WHERE id = 'eventy-002'"))  # NULL w firmie, tekst w lustrze
        self.assertEqual(self.database.diagnose().mirror_drift, 1)

    def test_two_adopters_at_once_create_each_company_once(self):
        self.database.upgrade_schema()
        reports, errors = [], []
        barrier = threading.Barrier(2)

        def worker():
            try:
                barrier.wait()
                reports.append(self.database.adopt())
            except Exception as exc:  # noqa: BLE001 — test zbiera błędy wątków
                errors.append(exc)

        threads = [threading.Thread(target=worker) for _ in range(2)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        self.assertEqual(errors, [])
        self.assertEqual(sorted(r.companies_added for r in reports), [0, 3])
        self.assertTrue(self.database.diagnose().clean)


@unittest.skipUnless(TEST_DB, SKIP_REASON)
class MysqlSiblingTests(unittest.TestCase):
    """Adresy-siostry na prawdziwym MySQL: indeks unikalny bez wielkości liter, podzapytania `hold` i `mail_count`,
    kolejność z `~` w id, usuwanie z renumeracją."""

    @classmethod
    def setUpClass(cls):
        cls.database = dbmod.Database(dbmod.make_engine(resolve_test_url()))
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)
        cls.database.ensure_schema()

    @classmethod
    def tearDownClass(cls):
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)

    def setUp(self):
        with self.database.transaction() as conn:
            conn.execute(dbmod.maile.delete())
            conn.execute(dbmod.companies.delete())
            self.database.insert_rows(conn, [
                sample_row(1, id="eventy-n001", row_number="", company="Beta", email="a@beta.pl"),
                sample_row(2, id="eventy-n0010", row_number="", company="Alfa", email="a@alfa.pl")])

    def test_add_hold_order_and_delete(self):
        self.database.add_mail("eventy-n001", "b@beta.pl")
        third = self.database.add_mail("eventy-n001", "c@beta.pl")
        self.assertEqual(self.database.list_ids(RowFilter(), "default"),
                         ["eventy-n001", "eventy-n001~2", "eventy-n001~3", "eventy-n0010"])
        self.assertEqual(self.database.list_ids(RowFilter(), "company"),
                         ["eventy-n0010", "eventy-n001", "eventy-n001~2", "eventy-n001~3"])
        self.database.apply_edit("eventy-n001", {"status": st.STATUS_REPLIED})
        self.assertEqual((self.database.get(third["id"])["hold"], self.database.get(third["id"])["mail_count"]),
                         (st.STATUS_REPLIED, "3"))
        self.assertEqual(self.database.get("eventy-n0010")["hold"], "")
        self.database.delete_mail("eventy-n001~2")
        self.assertEqual([(m["id"], m["position"]) for m in self.database.company_mails("eventy-n001")],
                         [("eventy-n001", "1"), ("eventy-n001~3", "2")])
        self.assertTrue(self.database.diagnose().clean, self.database.diagnose())

    def test_claim_has_exactly_one_winner_among_the_siblings_of_a_company(self):
        ids = ["eventy-n001", self.database.add_mail("eventy-n001", "b@beta.pl")["id"],
               self.database.add_mail("eventy-n001", "c@beta.pl")["id"]]
        winners, barrier = [], threading.Barrier(9)

        def worker(row_id):
            barrier.wait()
            if self.database.claim_for_sending(row_id):
                winners.append(row_id)

        threads = [threading.Thread(target=worker, args=(ids[n % 3],)) for n in range(9)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        self.assertEqual(len(winners), 1)
        self.assertEqual(self.database.count_rows(RowFilter(statuses=(st.STATUS_SENDING,))), 1)

    def test_claim_is_refused_while_a_sibling_holds_the_company(self):
        sibling = self.database.add_mail("eventy-n001", "b@beta.pl")["id"]
        self.database.apply_edit("eventy-n001", {"status": st.STATUS_DO_NOT_CONTACT})
        self.assertIsNone(self.database.claim_for_sending(sibling))
        self.assertEqual(self.database.list_ids(RowFilter(ready=True), "default"), ["eventy-n0010"])
        self.assertEqual([r["id"] for r in self.database.select_candidates(limit=None, unheld=True, per_company=True)],
                         ["eventy-n0010"])

    def test_grouped_sorts_filters_and_computed_columns_on_mysql(self):
        """Funkcje okienkowe z kolacją w ORDER BY, EXISTS dla filtrów i podzapytania wyliczanych kolumn."""
        self.database.add_mail("eventy-n001", "b@beta.pl")
        self.database.apply_edit("eventy-n001~2", {"status": st.STATUS_ERROR})
        self.database.apply_edit("eventy-n0010", {"status": st.STATUS_SENT})
        self.assertEqual(self.database.list_ids(RowFilter(), "status"), ["eventy-n001", "eventy-n0010", "eventy-n001~2"])
        self.assertEqual(self.database.list_ids(RowFilter(), "g:status"), ["eventy-n001", "eventy-n001~2", "eventy-n0010"])
        self.assertEqual(self.database.list_ids(RowFilter(), "g:-status"), ["eventy-n001", "eventy-n001~2", "eventy-n0010"])
        self.assertEqual(self.database.list_ids(RowFilter(), "g:email"), ["eventy-n0010", "eventy-n001", "eventy-n001~2"])
        self.assertEqual(self.database.list_ids(RowFilter(), "g:-sent_at")[:2], ["eventy-n0010", "eventy-n001"])
        self.assertEqual(self.database.list_ids(RowFilter(), "-mails"), ["eventy-n001", "eventy-n001~2", "eventy-n0010"])
        self.assertEqual(self.database.list_ids(RowFilter(multi=True), "default"), ["eventy-n001", "eventy-n001~2"])
        self.assertEqual(self.database.list_ids(RowFilter(multi=False), "default"), ["eventy-n0010"])
        self.database.apply_edit("eventy-n001", {"status": st.STATUS_REPLIED})
        self.assertEqual(self.database.list_ids(RowFilter(held=True), "default"), ["eventy-n001~2"])
        rows = {r["id"]: r for r in self.database.iter_rows(RowFilter(), ("id", "company_id", "position", "mail_count", "hold"))}
        self.assertEqual((rows["eventy-n001~2"]["mail_count"], rows["eventy-n001~2"]["hold"], rows["eventy-n001~2"]["position"]),
                         (2, st.STATUS_REPLIED, 2))

    def test_sibling_subqueries_use_the_company_index_whatever_the_statistics_say(self):
        """Zaraz po masowym wczytaniu statystyki MySQL bywają nieaktualne, a bez podpowiedzi optymalizator czyta wtedy całą tabelę
        dla każdego wiersza (50 tys. firm: minuty). Podpowiedź FORCE INDEX ma trzymać plan na indeksie (company_id, status)."""
        from sqlalchemy.dialects import mysql

        from perno.db.company_sql import select_mails
        from perno.db.listing_queries import light_columns
        from perno.db.queries import ordering
        from perno.db.schema import COMPANY_STATUS_INDEX

        self.database.add_mail("eventy-n001", "b@beta.pl")
        queries = {
            "light list": select_mails(*light_columns()).order_by(*ordering()),
            "ready": select_mails(dbmod.maile.c.id).where(*self.database._where(RowFilter(ready=True, multi=True, held=False))),
        }
        for name, statement in queries.items():
            sql = str(statement.compile(dialect=mysql.dialect(), compile_kwargs={"literal_binds": True}))
            with self.database.engine.connect() as conn:
                plan = [dict(row) for row in conn.execute(sa.text("EXPLAIN FORMAT=TRADITIONAL " + sql)).mappings()]
            siblings = [row for row in plan if str(row["table"]).startswith(("sibling", "earlier"))]
            self.assertTrue(siblings, name)
            for row in siblings:
                self.assertEqual(row["key"], COMPANY_STATUS_INDEX, (name, row))

    def test_queue_queries_for_the_company_gap(self):
        from perno.domain import automat_rules as ar

        self.database.add_mail("eventy-n001", "b@beta.pl")
        self.database.claim_for_sending("eventy-n001")
        self.database.mark_sent("eventy-n001")
        since = datetime(2000, 1, 1)
        self.assertEqual(self.database.last_sent_by_company(since), {"eventy-n001": datetime.now().date()})
        self.assertEqual(self.database.last_sent_by_company(datetime.now().replace(hour=0, minute=0, second=0) + timedelta(days=1)), {})
        settings = ar.DEFAULT_SETTINGS
        base = ar.queue_filter(settings)
        self.assertEqual(self.database.queue_company_ids(base, "default"), ["eventy-n001", "eventy-n0010"])
        self.assertEqual(self.database.count_companies(base), 2)
        self.assertEqual(self.database.count_matching(ar.queue_filter(settings, ("eventy-n001",))), 1)
        self.assertEqual(self.database.list_ids(ar.queue_filter(settings, ("eventy-n001",)), "default", limit=1), ["eventy-n0010"])

    def test_the_unique_index_is_case_insensitive_on_mysql(self):
        self.database.add_mail("eventy-n001", "b@beta.pl")
        with self.assertRaises(dbmod.InvalidValue):
            self.database.add_mail("eventy-n001", "B@BETA.PL")
        with self.assertRaises(sa_exc.IntegrityError), self.database.engine.begin() as conn:  # bez kontroli w kodzie
            conn.execute(sa.text("INSERT INTO maile (id, segment, email, status, retry_count, company_id, position) "
                                 "VALUES ('x', 'eventy', 'B@Beta.pl', 'do_wyslania', 0, 'eventy-n001', 9)"))


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(TEST_DB, SKIP_REASON)
class MysqlRepliesSmokeTests(unittest.TestCase):
    """Tabele i zapytania modułu odpowiedzi na prawdziwym MySQL: unikat skrótu, utf8mb4, długie treści, FOR UPDATE, blokada przebiegu."""

    @classmethod
    def setUpClass(cls):
        from perno.db.replies_schema import replies_metadata
        from perno.db.replies_store import ReplyStore

        cls.database = dbmod.Database(dbmod.make_engine(resolve_test_url()))
        replies_metadata.drop_all(cls.database.engine)
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)
        cls.database.ensure_schema()
        cls.store = ReplyStore(cls.database)
        cls.store.ensure_schema()

    @classmethod
    def tearDownClass(cls):
        from perno.db.replies_schema import replies_metadata

        replies_metadata.drop_all(cls.database.engine)
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)

    def setUp(self):
        from perno.db.replies_schema import replies_metadata

        with self.database.transaction() as conn:
            for table in reversed(replies_metadata.sorted_tables):
                if table.name != "odp_przebieg":
                    conn.execute(table.delete())
            conn.execute(sa.text("UPDATE odp_przebieg SET token = NULL, running_since = NULL"))
            conn.execute(dbmod.maile.delete())
            conn.execute(dbmod.companies.delete())

    def message(self, identity="a@x.pl", **overrides):
        from perno.domain import replies as rp

        values = dict(direction="in", kind="human", mailbox="roman@z3x.pl", identity=identity, message_id=identity, thread_id="eventy-001",
                      company_id="eventy-001", date_at=datetime(2026, 10, 8, 10, 0), subject="Temat", from_addr="a@b.pl", body_text="Treść",
                      match_method="address", thread_mailbox="roman@z3x.pl")
        values.update(overrides)
        return rp.NewMessage(**values)

    def test_a_message_round_trips_with_emoji_polish_letters_and_a_very_long_body(self):
        body = "Zażółć gęślą jaźń 😀 " + "x" * 200_000
        stored = self.store.save_message(self.message(body_text=body, subject="Re: Wiadomość 📧 ąęść"), datetime(2026, 10, 8, 10, 1))
        self.assertIsNotNone(stored)
        with self.database.engine.connect() as conn:
            row = conn.execute(sa.text("SELECT subject, CHAR_LENGTH(body_text), body_text FROM odp_wiadomosci")).first()
        self.assertEqual((row[0], row[1]), ("Re: Wiadomość 📧 ąęść", len(body)))
        self.assertTrue(row[2].startswith("Zażółć gęślą jaźń 😀"))

    def test_the_unique_hash_stops_duplicates_even_for_a_message_id_longer_than_an_index_would_allow(self):
        long_id = "z" * 900 + "@x.pl"
        self.assertIsNotNone(self.store.save_message(self.message(long_id), datetime(2026, 10, 8, 10, 1)))
        self.assertIsNone(self.store.save_message(self.message(long_id), datetime(2026, 10, 8, 10, 2)))
        self.assertEqual(self.store.counts()["messages"], 1)
        self.assertEqual(self.store.known_threads([long_id]), {long_id: "eventy-001"})

    def test_the_hash_column_is_case_and_accent_exact_and_ascii(self):
        with self.database.engine.connect() as conn:
            row = conn.execute(sa.text(
                "SELECT CHARACTER_SET_NAME, COLLATION_NAME FROM information_schema.COLUMNS "
                "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = 'odp_wiadomosci' AND COLUMN_NAME = 'message_id_hash'")).first()
        self.assertEqual((row[0], row[1]), ("ascii", "ascii_bin"))

    def test_the_thread_is_created_and_updated_under_for_update(self):
        put_rows(self.database, sample_row(1, status="wyslano", sent_at="2026-09-29T13:00:00"))  # wątek istnieje tylko razem z mailem kampanii
        for n in range(3):
            self.store.save_message(self.message(f"m{n}@x.pl", date_at=datetime(2026, 10, 8, 10, n)), datetime(2026, 10, 8, 11, 0))
        counts = self.store.counts()
        self.assertEqual((counts["threads"], counts["awaiting_reply"], counts["unread"]), (1, 1, 1))

    def test_the_run_lock_is_exclusive_between_threads(self):
        now = datetime(2026, 10, 8, 10, 0)
        results = []
        barrier = threading.Barrier(6)

        def attempt(token):
            barrier.wait()
            results.append(self.store.acquire_run(token, now))

        threads = [threading.Thread(target=attempt, args=(f"t{n}",)) for n in range(6)]
        [t.start() for t in threads]
        [t.join() for t in threads]
        self.assertEqual(sorted(results), [False] * 5 + [True])

    def test_the_cursor_with_a_composite_key_and_the_state_summary(self):
        from perno.db.replies_store import Cursor

        now = datetime(2026, 10, 8, 10, 0)
        self.store.save_cursor("roman@z3x.pl", "Wysłane", Cursor(4_000_000_000, 5_000_000_000), now)  # powyżej 32 bitów
        self.store.record_error("roman@z3x.pl", "błąd łączenia ✗", "network", now)
        self.assertEqual(self.store.get_cursor("roman@z3x.pl", "Wysłane"), Cursor(4_000_000_000, 5_000_000_000))
        self.assertEqual(self.store.mailbox_states()[0]["error"], "błąd łączenia ✗")

    def test_sent_mails_and_the_status_transition_work_with_the_company_join(self):
        with self.database.transaction() as conn:
            self.database.insert_rows(conn, [sample_row(1, status="wyslano", sent_at="2026-09-29T13:00:00", website="https://mylimo.pl")])
        self.database.mark_sent("eventy-001", "roman@z3x.pl", "orig1@z3x.pl")
        rows = self.store.sent_mails()
        self.assertEqual([(r["id"], r["message_id"], r["website"]) for r in rows], [("eventy-001", "orig1@z3x.pl", "https://mylimo.pl")])
        self.assertTrue(self.database.transition_status("eventy-001", ["wyslano"], "odpowiedzial"))
        self.assertFalse(self.database.transition_status("eventy-001", ["wyslano"], "zwrot"))
        self.assertEqual(self.database.get("eventy-001")["status"], "odpowiedzial")

    def test_the_list_queries_run_on_mysql_for_every_tab_filter_and_search(self):
        from perno.services import replies as svc

        with self.database.transaction() as conn:
            self.database.insert_rows(conn, [sample_row(n, company=f"Firma ąć {n}", country="PL", status="wyslano", sent_at="2026-09-29T13:00:00")
                                             for n in (1, 2, 3)])
        for n, kind in ((1, "human"), (2, "bounce"), (3, "human")):
            self.store.save_message(self.message(f"l{n}@x.pl", thread_id=f"eventy-00{n}", company_id=f"eventy-00{n}", kind=kind,
                                                 body_text="Zażółć 100% pewne _tak_ \\ ukośnik", snippet="Zażółć gęślą"), datetime(2026, 10, 8, 11, 0))
        self.store.set_rating("eventy-003", "positive", "human", "", datetime(2026, 10, 8, 12, 0))
        for tab in ("todo", "all", "positive", "neutral", "negative", "unrated", "machine", "unassigned"):
            result = svc.list_replies(self.store, {"tab": tab, "mailboxes": ["Roman@z3x.pl"], "segments": ["eventy"], "countries": ["PL"],
                                                   "from": "2026-10-01", "to": "2026-10-31", "q": "firma"})
            self.assertEqual(result["tab"], tab)
        self.assertEqual([r["id"] for r in svc.list_replies(self.store, {"tab": "all"})["rows"]], ["eventy-001", "eventy-003"])  # ta sama chwila: po id
        self.assertEqual(svc.list_replies(self.store, {"tab": "positive"})["counts"]["positive"], 1)
        for query, expected in (("100%", 2), ("_tak_", 2), ("ukośnik", 2), ("ZAŻÓŁĆ", 2), ("zażolc", 2), ("100% pewne", 2), ("%%", 0), ("\\", 2)):
            found = svc.list_replies(self.store, {"tab": "all", "q": query})["total"]
            self.assertEqual(found, expected, query)
        self.assertEqual(svc.list_replies(self.store, {"tab": "machine"})["total"], 1)
        view = svc.thread_view(self.store, "eventy-001")
        self.assertEqual((view["thread"]["company"], len(view["messages"])), ("Firma ąć 1", 1))

    def test_the_report_queries_run_on_mysql(self):
        from perno.db.schema import maile
        from perno.services import replies_report as rep

        with self.database.transaction() as conn:
            self.database.insert_rows(conn, [sample_row(n, company=f"Firma {n}", country="PL" if n < 3 else "DE", status="odpowiedzial" if n < 3 else "zwrot",
                                                        sent_at="2026-09-29T13:00:00") for n in (1, 2, 3)])
            conn.execute(maile.update().where(maile.c.id == "eventy-002").values(sent_from="drugi@z3x.pl"))
        for n, kind in ((1, "human"), (2, "human"), (3, "bounce")):
            self.store.save_message(self.message(f"r{n}@x.pl", thread_id=f"eventy-00{n}", company_id=f"eventy-00{n}", kind=kind,
                                                 date_at=datetime(2026, 9, 30, 9, 0), bounce_type="hard" if kind == "bounce" else ""), datetime(2026, 10, 8, 11, 0))
        self.store.set_rating("eventy-001", "positive", "human", "", datetime(2026, 10, 8, 12, 0))
        config = {"SMTP_USER": "roman@z3x.pl", "SMTP_HOST": "s", "SMTP_PORT": "1", "SMTP_PASSWORD": "p"}
        report = rep.build_report(self.store, config, {"from": "2026-09-01", "to": "2026-10-31", "segments": ["eventy"], "mailboxes": ["roman@z3x.pl", "drugi@z3x.pl"]},
                                  datetime(2026, 10, 8, 12, 0))
        self.assertEqual((report["total"]["sent"], report["total"]["bounced"], report["total"]["replied"], report["total"]["positive"]), (3, 1, 2, 1))
        self.assertEqual({row["key"]: row["sent"] for row in report["groups"]["mailbox"]}, {"roman@z3x.pl": 2, "drugi@z3x.pl": 1})
        self.assertEqual([row["key"] for row in report["groups"]["day"]], ["2026-09-29"])
        self.assertEqual((report["latency"]["n"], report["replied_companies"]), (2, 2))
        self.assertIn("Raport odpowiedzi", rep.markdown(report))


@unittest.skipUnless(TEST_DB, SKIP_REASON)
class MysqlLeadsSmokeTests(unittest.TestCase):
    """Leady na prawdziwym MySQL: zasiew z grupowaniem wątków po firmie, karty, sortowanie i wyszukiwanie (znaki LIKE), utf8mb4,
    blokada wiersza przy równoczesnym zapisie AI i człowieka oraz kopia tabeli posortowana po kluczu głównym."""

    @classmethod
    def setUpClass(cls):
        from perno.db.leads_store import LeadStore
        from perno.db.replies_schema import replies_metadata
        from perno.db.replies_store import ReplyStore

        cls.database = dbmod.Database(dbmod.make_engine(resolve_test_url()))
        replies_metadata.drop_all(cls.database.engine)
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)
        cls.database.ensure_schema()
        cls.replies = ReplyStore(cls.database)
        cls.replies.ensure_schema()
        cls.leads = LeadStore(cls.database)
        cls.leads.ensure_schema()

    @classmethod
    def tearDownClass(cls):
        from perno.db.replies_schema import replies_metadata

        replies_metadata.drop_all(cls.database.engine)
        dbmod.metadata.drop_all(cls.database.engine, tables=[dbmod.maile, dbmod.companies], checkfirst=True)

    def setUp(self):
        from perno.db.replies_schema import replies_metadata

        with self.database.transaction() as conn:
            for table in reversed(replies_metadata.sorted_tables):
                if table.name != "odp_przebieg":
                    conn.execute(table.delete())
            conn.execute(dbmod.maile.delete())
            conn.execute(dbmod.companies.delete())
        with self.database.transaction() as conn:
            self.database.insert_rows(conn, [sample_row(n, company=f"Firma ąć {n}", website=f"https://firma{n}.pl", country="PL", status="odpowiedzial",
                                                        sent_at="2026-09-29T13:00:00") for n in (1, 2, 3, 4)])

    def reply(self, n, rating, thread=None, when=datetime(2026, 10, 8, 10, 0), who="Anna Nowak"):
        from perno.domain import replies as rp

        thread = thread or f"eventy-00{n}"
        self.replies.save_message(rp.NewMessage(direction="in", kind="human", mailbox="roman@z3x.pl", identity=f"{thread}@x", message_id=f"{thread}@x",
                                                thread_id=thread, company_id=f"eventy-00{n}", date_at=when, subject="Re", from_addr=f"a{n}@firma{n}.pl",
                                                from_name=who, body_text="Poproszę ofertę", snippet="Poproszę ofertę", match_method="address",
                                                thread_mailbox="roman@z3x.pl"), when)
        self.replies.set_rating(thread, rating, "human", "", when)

    def test_leads_are_grouped_by_company_and_cards_sorting_and_search_run_on_mysql(self):
        from perno.domain import lead_listing as ll

        now = datetime(2026, 10, 9, 10, 0)
        self.database.add_mail("eventy-001", "drugi@firma1.pl")
        self.reply(1, "neutral")
        self.reply(1, "positive", thread="eventy-001~2", when=datetime(2026, 10, 8, 12, 0), who="Jan Szef")
        self.reply(2, "neutral", when=datetime(2026, 10, 8, 11, 0), who="Bartek")
        self.reply(3, "negative")
        self.assertEqual(self.leads.pending_total(), 2)
        self.assertEqual(self.leads.sync(now), 2)
        self.assertEqual(self.leads.sync(now), 0)
        first = self.leads.get("eventy-001")
        self.assertEqual((first["rating"], first["thread_id"], first["email"], first["contact"]), ("positive", "eventy-001~2", "a1@firma1.pl", "Jan Szef"))
        self.leads.update_fields("eventy-002", {"about": "Zażółć gęślą jaźń 😀 100% _pewne_"}, by="human", now=now)

        def ids(**kw):
            return [row["id"] for row in self.leads.list_page(ll.LeadFilter(**kw), 1, 10)[0]]

        self.assertEqual(self.leads.tab_counts(ll.LeadFilter()), {"all": 2, "missing": 1, "ai": 0, "withdrawn": 0})
        self.assertEqual((ids(), ids(sort=ll.SORT_NAME), ids(sort=ll.SORT_RATING), ids(sort=ll.SORT_COMPLETE)),
                         (["eventy-001", "eventy-002"], ["eventy-001", "eventy-002"], ["eventy-001", "eventy-002"], ["eventy-001", "eventy-002"]))
        self.assertEqual((ids(tab=ll.TAB_MISSING), ids(words=("100%",)), ids(words=("_pewne_",)), ids(words=("%%",)), ids(words=("ZAŻÓŁĆ",))),
                         (["eventy-001"], ["eventy-002"], ["eventy-002"], [], ["eventy-002"]))
        self.assertEqual(self.leads.get("eventy-002")["about"], "Zażółć gęślą jaźń 😀 100% _pewne_")
        self.replies.set_rating("eventy-001", "negative", "human", "", now)
        self.replies.set_rating("eventy-001~2", "negative", "human", "", now)
        self.assertEqual((ids(tab=ll.TAB_WITHDRAWN), self.leads.get("eventy-001")["active"]), (["eventy-001"], False))

    def test_a_human_and_the_ai_writing_the_same_field_at_once_never_end_with_the_ai_value(self):
        self.reply(1, "positive")
        now = datetime(2026, 10, 9, 10, 0)
        self.leads.sync(now)
        for round_no in range(8):
            self.leads.update_fields("eventy-001", {"about": ""}, by="human", now=now)
            barrier = threading.Barrier(2)

            def write(by, value, gate=barrier):
                gate.wait()
                self.leads.update_fields("eventy-001", {"about": value}, by=by, now=now)

            threads = [threading.Thread(target=write, args=("human", "Człowiek")), threading.Thread(target=write, args=("ai", "AI"))]
            [t.start() for t in threads]
            [t.join() for t in threads]
            lead = self.leads.get("eventy-001")
            self.assertEqual((lead["about"], lead["src"]), ("Człowiek", {"about": "human"}), round_no)

    def test_the_lead_table_is_dumped_in_primary_key_order_for_the_backup(self):
        import tempfile

        from perno.services import firms_import

        self.reply(1, "positive")
        self.reply(2, "positive")
        now = datetime(2026, 10, 9, 10, 0)
        self.leads.sync(now)
        with tempfile.TemporaryDirectory() as directory:
            firms_import.backup_table(self.database, directory, "test", replies=True)
            import os

            name = next(item for item in os.listdir(directory) if item.startswith("OdpLeady_"))
            with open(os.path.join(directory, name), encoding="utf-8") as handle:
                lines = handle.read().splitlines()
        self.assertEqual([line.split(",")[0] for line in lines[1:]], ["eventy-001", "eventy-002"])
