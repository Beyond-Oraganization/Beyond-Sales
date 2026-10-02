from __future__ import annotations

import csv
import io
import json
import tracemalloc
import unittest
from datetime import datetime

import openpyxl

from perno import db as dbmod
from perno.db import RowFilter
from perno.domain import statuses as st
from perno.services import export as exporter
from perno.web.app import App, Request
from perno.web.pages import PAGES
from tests.fixtures import make_db, put, sample_row

PORT = 8799
BODY = 'Linia 1, "cytat"; średnik\nLinia 2\r\nLinia 3'


def read_all(result) -> bytes:
    with result.stream as f:
        return f.read()


def csv_rows(raw: bytes, delimiter=";"):
    return list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=delimiter))


KEY_OF = {label: key for key, (label, _) in exporter.COLUMN_INFO.items()}
ALL = exporter.ALL_EXPORT_COLUMNS


def keyed(rows):
    """Wiersze tabeli jako słowniki o kluczach pól (nagłówki pliku są po polsku, więc tłumaczymy je z powrotem)."""
    header = [KEY_OF[label] for label in rows[0]]
    return [dict(zip(header, row)) for row in rows[1:]]


class FormatCase(unittest.TestCase):
    def setUp(self):
        self.database = make_db()
        put(self.database,
            sample_row(1, status="wyslano", sent_at="2026-09-29T14:23:11", company="Zażółć gęślą jaźń",
                       body_html="<p>html</p>\n<p>druga</p>", body_text=BODY, retry_count="2"),
            sample_row(2, segment="driverzy", subject="=1+1", company="+48 600 000 000", error="@cmd"),
            sample_row(3, contact_person="-ostrożnie", body_text="ctrl\x01znak"),
            sample_row(4, status="pominieto_klient", email="", body_html="", body_text=""))

    def export(self, fmt, **kw):
        spec = exporter.ExportSpec(fmt=fmt, **kw)
        return read_all(exporter.build_export(self.database, spec, now=datetime(2026, 9, 30, 15, 30)))


class ReadableFileTests(FormatCase):
    """Plik ma być czytelny: mało kolumn, polskie nagłówki, etykiety zamiast kodów."""

    def test_default_file_has_a_few_readable_columns_with_polish_headers(self):
        rows = csv_rows(self.export("csv"))
        self.assertEqual(rows[0], ["Firma", "Segment", "Kraj", "Miasto", "E-mail", "Strona www", "Status", "Data wysłania", "Temat maila"])
        self.assertLessEqual(len(rows[0]), 10)
        for technical in ("id", "row_number", "tag", "mode", "retry_count", "updated_at", "body_html"):
            self.assertNotIn(technical, rows[0])

    def test_status_is_a_label_and_dates_have_no_seconds(self):
        by_company = {r["company"]: r for r in keyed(csv_rows(self.export("csv")))}
        first = by_company["Zażółć gęślą jaźń"]
        self.assertEqual((first["status"], first["sent_at"]), ("Wysłano", "2026-09-29 14:23"))
        self.assertEqual(by_company["Firma 3"]["status"], "Do wysłania")
        self.assertEqual(by_company["Firma 3"]["sent_at"], "")
        self.assertEqual(next(r for r in keyed(csv_rows(self.export("csv"))) if r["segment"] == "eventy" and r["email"] == "")["status"], "Już klient")

    def test_unknown_status_code_is_shown_as_is(self):
        self.assertEqual(exporter.status_label("cos_nowego"), "cos_nowego")
        self.assertEqual(exporter.status_label(None), "")

    def test_columns_are_ordered_for_people_with_technical_ones_last(self):
        order = exporter.ALL_EXPORT_COLUMNS
        self.assertEqual(order[:3], ("company", "segment", "country_name"))
        self.assertLess(order.index("status"), order.index("body_html"))
        self.assertLess(order.index("body_text"), order.index("id"))
        self.assertEqual(order[-1], "updated_at")
        self.assertEqual(len(order), len(set(order)))

    def test_every_column_has_a_distinct_polish_header_and_a_width(self):
        labels = [exporter.COLUMN_INFO[name][0] for name in exporter.ALL_EXPORT_COLUMNS]
        self.assertEqual(len(labels), len(set(labels)))
        for name in exporter.ALL_EXPORT_COLUMNS:
            self.assertIn(name, exporter._WIDTHS)
        self.assertEqual(exporter.headers(("company", "sent_at")), ("Firma", "Data wysłania"))

    def test_catalog_covers_exactly_the_database_columns_plus_the_derived_one(self):
        self.assertEqual(set(exporter.ALL_EXPORT_COLUMNS), set(dbmod.COLUMNS) | {"updated_at", "country_name"})

    def test_presets_are_ordered_from_compact_to_everything(self):
        sizes = {key: len(columns) for key, _, columns in exporter.COLUMN_PRESETS}
        self.assertLess(sizes["minimal"], sizes["contacts"])
        self.assertLessEqual(sizes["contacts"], sizes["standard"])
        self.assertLess(sizes["standard"], sizes["extended"])
        self.assertLess(sizes["extended"], sizes["all"])
        self.assertEqual(sizes["all"], len(exporter.ALL_EXPORT_COLUMNS))


class CsvTests(FormatCase):
    def test_bom_once_and_default_columns(self):
        raw = self.export("csv")
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(raw.count(b"\xef\xbb\xbf"), 1)
        rows = csv_rows(raw)
        self.assertEqual(tuple(rows[0]), exporter.headers(exporter.EXPORT_COLUMNS))
        self.assertEqual(len(rows), 5)

    def test_values_and_polish_characters(self):
        first = keyed(csv_rows(self.export("csv")))[0]
        self.assertEqual(first["company"], "Zażółć gęślą jaźń")
        self.assertEqual(first["sent_at"], "2026-09-29 14:23")
        self.assertEqual(keyed(csv_rows(self.export("csv")))[1]["sent_at"], "")

    def test_technical_columns_on_request(self):
        first = keyed(csv_rows(self.export("csv", columns=("id", "retry_count", "updated_at"))))[0]
        self.assertEqual((first["id"], first["retry_count"]), ("eventy-001", "2"))
        self.assertRegex(first["updated_at"], r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$")

    def test_formula_injection_guard(self):
        rows = {r["id"]: r for r in keyed(csv_rows(self.export("csv", columns=ALL)))}
        self.assertEqual(rows["driverzy-002"]["subject"], "'=1+1")
        self.assertEqual(rows["driverzy-002"]["company"], "'+48 600 000 000")
        self.assertEqual(rows["driverzy-002"]["error"], "'@cmd")
        self.assertEqual(rows["eventy-003"]["contact_person"], "'-ostrożnie")
        self.assertEqual(rows["eventy-001"]["id"], "eventy-001")

    def test_delimiters(self):
        n = len(exporter.EXPORT_COLUMNS)
        self.assertEqual(len(csv_rows(self.export("csv", delimiter=","), ",")[0]), n)
        semicolon = self.export("csv", delimiter=";")
        self.assertEqual(len(csv_rows(semicolon, ";")[0]), n)
        self.assertEqual(len(csv_rows(semicolon, ",")[0]), 1)

    def test_bodies_toggle_adds_the_message_columns_and_keeps_multiline_text_intact(self):
        rows = csv_rows(self.export("csv", bodies=True))
        wanted = [c for c in exporter.ALL_EXPORT_COLUMNS if c in exporter.EXPORT_COLUMNS or c in exporter.BODY_COLUMNS]
        self.assertEqual(tuple(rows[0]), exporter.headers(tuple(wanted)))
        first = keyed(rows)[0]
        self.assertEqual(first["body_text"], BODY)
        self.assertEqual(first["body_html"], "<p>html</p>\n<p>druga</p>")
        self.assertNotIn("id", first)

    def test_filters_and_empty_result(self):
        flt = RowFilter(segments=("eventy",), statuses=("do_wyslania",))
        rows = keyed(csv_rows(self.export("csv", flt=flt)))
        self.assertEqual([r["company"] for r in rows], ["Firma 3"])
        empty = csv_rows(self.export("csv", flt=RowFilter(statuses=("blad",))))
        self.assertEqual(len(empty), 1)


class CountryColumnsTests(unittest.TestCase):
    def setUp(self):
        self.database = make_db()
        put(self.database,
            sample_row(1, country="PL", city="Kraków", address="ul. Długa 5", description="Kursy survivalu"),
            sample_row(2, segment="driverzy", country="UK", city="Londyn"),
            sample_row(3))

    def export(self, fmt, **kw):
        spec = exporter.ExportSpec(fmt=fmt, **kw)
        return read_all(exporter.build_export(self.database, spec, now=datetime(2026, 9, 30, 15, 30)))

    def test_country_is_shown_by_name_in_the_default_file_and_the_iso_code_is_optional(self):
        self.assertIn("country_name", exporter.EXPORT_COLUMNS)
        self.assertNotIn("country", exporter.EXPORT_COLUMNS)
        self.assertNotIn("address", exporter.EXPORT_COLUMNS)  # rzadko wypełniony i długi — do dobrania w „Rozszerzonym”
        order = exporter.ALL_EXPORT_COLUMNS
        self.assertEqual(order[order.index("country_name"):order.index("country_name") + 4], ("country_name", "country", "city", "address"))

    def test_csv_has_name_and_optional_iso_code(self):
        rows = keyed(csv_rows(self.export("csv", columns=("company", "country_name", "country", "city", "address", "description"))))
        by_company = {r["company"]: r for r in rows}
        first, second, third = by_company["Firma 1"], by_company["Firma 2"], by_company["Firma 3"]
        self.assertEqual((first["country"], first["country_name"], first["city"], first["address"], first["description"]),
                         ("PL", "Polska", "Kraków", "ul. Długa 5", "Kursy survivalu"))
        self.assertEqual((second["country"], second["country_name"]), ("GB", "Wielka Brytania"))  # UK trafia do bazy jako GB
        self.assertEqual((third["country"], third["country_name"], third["city"]), ("", "", ""))

    def test_xlsx_has_the_same_columns_and_values(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.export("xlsx", columns=("company", "country_name", "country", "city"))))
        ws = wb["Maile"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header, ["Firma", "Kraj", "Kod kraju", "Miasto"])
        self.assertEqual([c.value for c in ws[2]], ["Firma 1", "Polska", "PL", "Kraków"])

    def test_bodies_option_keeps_the_country_column(self):
        rows = csv_rows(self.export("csv", bodies=True))
        self.assertIn("Kraj", rows[0])
        self.assertIn("Treść maila (HTML)", rows[0])

    def test_city_and_address_cells_are_protected_like_any_other_text(self):
        put(self.database, sample_row(9, city="=HYPERLINK(\"http://x\")", address="@cmd"))
        rows = {r["company"]: r for r in keyed(csv_rows(self.export("csv", columns=("company", "city", "address"))))}
        self.assertEqual(rows["Firma 9"]["city"], "'=HYPERLINK(\"http://x\")")
        self.assertEqual(rows["Firma 9"]["address"], "'@cmd")


class XlsxTests(FormatCase):
    def load(self, **kw):
        wb = openpyxl.load_workbook(io.BytesIO(self.export("xlsx", **kw)))
        return wb, wb["Maile"]

    def cells(self, ws, row, header=None):
        header = header or [KEY_OF[c.value] for c in ws[1]]
        return {name: ws.cell(row=row, column=i + 1) for i, name in enumerate(header)}

    def test_sheet_layout(self):
        wb, ws = self.load()
        self.assertEqual(wb.sheetnames, ["Maile"])
        self.assertEqual([c.value for c in ws[1]], list(exporter.headers(exporter.EXPORT_COLUMNS)))
        self.assertTrue(all(c.font.bold for c in ws[1]))
        self.assertEqual(ws.freeze_panes, "B2")  # nagłówek i nazwa firmy zostają na ekranie
        self.assertEqual(ws.auto_filter.ref, f"A1:{openpyxl.utils.get_column_letter(len(exporter.EXPORT_COLUMNS))}5")
        self.assertEqual(ws.max_row, 5)

    def test_header_is_styled_and_columns_have_widths(self):
        _, ws = self.load()
        cell = ws["A1"]
        self.assertEqual(cell.font.color.rgb, "FFFFFFFF")
        self.assertEqual(cell.fill.start_color.rgb, "FF1F3A5F")
        self.assertGreater(ws.row_dimensions[1].height, 18)
        self.assertEqual(ws.column_dimensions["A"].width, exporter._WIDTHS["company"])
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(ws.print_title_rows, "$1:$1")

    def test_body_cells_wrap_and_align_to_top(self):
        _, ws = self.load(columns=("company", "email", "status", "sent_at", "retry_count", "subject"))
        for row in (2, 3):
            for name, cell in self.cells(ws, row).items():
                if cell.value is None:
                    continue
                self.assertEqual(cell.alignment.vertical, "top", name)
                self.assertTrue(cell.alignment.wrap_text, name)

    def test_freeze_panes_without_the_company_column(self):
        _, ws = self.load(columns=("email", "status"))
        self.assertEqual(ws.freeze_panes, "A2")

    def test_typed_cells(self):
        _, ws = self.load(columns=("company", "sent_at", "retry_count"))
        first = self.cells(ws, 2)
        self.assertEqual(first["sent_at"].value, datetime(2026, 9, 29, 14, 23, 11))
        self.assertEqual(first["sent_at"].number_format, exporter.DATETIME_XLSX)
        self.assertEqual(first["retry_count"].value, 2)
        self.assertEqual(first["company"].value, "Zażółć gęślą jaźń")
        self.assertIsNone(self.cells(ws, 3)["sent_at"].value)

    def test_status_is_a_coloured_label(self):
        _, ws = self.load(columns=("company", "status"))
        sent, pending = self.cells(ws, 2)["status"], self.cells(ws, 3)["status"]
        self.assertEqual((sent.value, pending.value), ("Wysłano", "Do wysłania"))
        self.assertEqual(sent.font.color.rgb, "FF1A7F4B")
        self.assertEqual(pending.font.color.rgb, "FF2F6FEB")
        self.assertTrue(sent.font.b)

    def test_website_is_a_clickable_link_but_only_for_http_urls(self):
        put(self.database, sample_row(7, website="javascript:alert(1)", company="Zły link"))
        _, ws = self.load(columns=("company", "website"))
        rows = {self.cells(ws, r)["company"].value: self.cells(ws, r)["website"] for r in range(2, ws.max_row + 1)}
        self.assertEqual(rows["Zażółć gęślą jaźń"].hyperlink.target, "https://firma1.example.com/")
        self.assertEqual(rows["Zażółć gęślą jaźń"].font.underline, "single")
        self.assertIsNone(rows["Zły link"].hyperlink)
        self.assertEqual(rows["Zły link"].value, "javascript:alert(1)")

    def test_formula_strings_stay_text(self):
        _, ws = self.load(columns=("company", "subject"))
        cell = self.cells(ws, 5)["subject"]
        self.assertEqual((cell.value, cell.data_type), ("=1+1", "s"))
        plus = self.cells(ws, 5)["company"]
        self.assertEqual((plus.value, plus.data_type), ("+48 600 000 000", "s"))

    def test_illegal_characters_are_stripped_and_bodies_included(self):
        _, ws = self.load(bodies=True)
        self.assertEqual(self.cells(ws, 3)["body_text"].value, "ctrlznak")
        self.assertEqual(self.cells(ws, 2)["body_html"].value, "<p>html</p>\n<p>druga</p>")

    def test_empty_result_is_a_valid_workbook(self):
        _, ws = self.load(flt=RowFilter(statuses=("blad",)))
        self.assertEqual(ws.max_row, 1)
        self.assertEqual(ws.auto_filter.ref, f"A1:{openpyxl.utils.get_column_letter(len(exporter.EXPORT_COLUMNS))}1")

    def test_filename_and_unknown_format(self):
        result = exporter.build_export(self.database, exporter.ExportSpec(fmt="xlsx"), now=datetime(2026, 9, 30, 15, 30))
        result.stream.close()
        self.assertEqual(result.filename, "Maile_2026-09-30_1530.xlsx")
        with self.assertRaises(exporter.ExportError):
            exporter.build_export(self.database, exporter.ExportSpec(fmt="pdf"))


class LargeTableTests(unittest.TestCase):
    ROWS = 20_000

    @classmethod
    def setUpClass(cls):
        cls.database = make_db()
        rows = [sample_row(n % 999 + 1, segment=f"seg{n}", company=f"Firma {n}", body_html="<p>x</p>" * 50,
                           body_text="tekst " * 50) for n in range(cls.ROWS)]
        with cls.database.transaction() as conn:
            cls.database.insert_rows(conn, rows)

    def measure(self, fmt):
        tracemalloc.start()
        try:
            result = exporter.build_export(self.database, exporter.ExportSpec(fmt=fmt, bodies=True))
            _, peak = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
        return result, peak

    def test_csv_streams_without_holding_everything_in_memory(self):
        result, peak = self.measure("csv")
        with result.stream as f, io.TextIOWrapper(f, encoding="utf-8-sig", newline="") as text:
            self.assertEqual(sum(1 for _ in csv.reader(text)), self.ROWS + 1)
        self.assertLess(peak, 5 * 1024 * 1024, f"szczyt pamięci: {peak / 1e6:.1f} MB")
        self.assertGreater(result.length, 10_000_000)

    def test_xlsx_streams_too(self):
        result, peak = self.measure("xlsx")
        with result.stream as f:
            wb = openpyxl.load_workbook(f, read_only=True)
            self.assertEqual(sum(1 for _ in wb["Maile"].iter_rows()), self.ROWS + 1)
        self.assertLess(peak, 30 * 1024 * 1024, f"szczyt pamięci: {peak / 1e6:.1f} MB")


class EndpointTests(unittest.TestCase):
    def setUp(self):
        self.database = make_db()
        put(self.database,
            sample_row(1, status="wyslano", sent_at="2026-09-29T14:23:11"),
            sample_row(2), sample_row(3, segment="driverzy"))
        self.app = App(self.database, PAGES, index_html=None, allowed_hosts={f"localhost:{PORT}"})

    def get(self, path, **query):
        return self.app.dispatch(Request("GET", path, query={k: v if isinstance(v, list) else [v] for k, v in query.items()},
                                         headers={"host": f"localhost:{PORT}"}))

    def test_export_page_is_in_the_menu(self):
        html = self.get("/eksport").body.decode("utf-8")
        self.assertIn('href="/eksport" aria-current="page">Eksport</a>', html)
        self.assertIn('href="/">Wysyłka</a>', html)
        self.assertLess(self.get("/").body.decode("utf-8").index(">Wysyłka<"),
                        self.get("/").body.decode("utf-8").index(">Eksport<"))

    def test_meta_and_count(self):
        meta = json.loads(self.get("/api/export/meta").body)
        self.assertEqual(meta["segments"], ["eventy", "driverzy"])
        self.assertEqual(meta["total"], 3)
        codes = [s["code"] for s in meta["statuses"]]
        self.assertEqual(codes, list(st.STATUS_META))
        self.assertEqual({s["code"]: s["count"] for s in meta["statuses"]}["wyslano"], 1)
        self.assertEqual(json.loads(self.get("/api/export/count").body), {"count": 3, "total": 3})
        filtered = self.get("/api/export/count", segment="eventy", status=["do_wyslania", "blad"])
        self.assertEqual(json.loads(filtered.body), {"count": 1, "total": 3})
        self.assertEqual(json.loads(self.get("/api/export/count", segment="").body)["count"], 3)

    def test_parameter_validation(self):
        self.assertEqual(self.get("/api/export/count", segment="nie-ma").status, 400)
        self.assertEqual(self.get("/api/export/count", status="nie-ma").status, 400)
        self.assertEqual(self.get("/api/export").status, 400)
        self.assertEqual(self.get("/api/export", format="pdf").status, 400)
        self.assertEqual(self.get("/api/export", format="csv", delimiter="tab").status, 400)

    def test_download_headers_and_content_match_count(self):
        resp = self.get("/api/export", format="csv", segment="eventy")
        self.assertEqual(resp.status, 200)
        self.assertRegex(resp.headers["Content-Disposition"], r'^attachment; filename="Maile_\d{4}-\d{2}-\d{2}_\d{4}\.csv"$')
        self.assertTrue(resp.content_type.startswith("text/csv"))
        with resp.stream as f:
            raw = f.read()
        self.assertEqual(len(raw), resp.length)
        self.assertEqual(len(csv_rows(raw)) - 1, 2)
        xlsx = self.get("/api/export", format="xlsx")
        self.assertIn("spreadsheetml", xlsx.content_type)
        self.assertTrue(xlsx.headers["Content-Disposition"].endswith('.xlsx"'))
        xlsx.stream.close()


class SelectionTests(unittest.TestCase):
    """Eksport wskazanych wierszy i wybranych kolumn (`RowFilter.ids`, `ExportSpec.columns`)."""

    def setUp(self):
        self.database = make_db()
        put(self.database,
            sample_row(1, country="GB", city="London", company="Alfa"),
            sample_row(2, segment="driverzy", country="PL", company="Beta"),
            sample_row(3, company="Gamma"))

    def build(self, fmt="csv", ids=None, **kwargs):
        spec = exporter.ExportSpec(fmt=fmt, flt=RowFilter(ids=ids), **kwargs)
        return exporter.build_export(self.database, spec, now=datetime(2026, 9, 30, 15, 30))

    def table(self, result, fmt="csv"):
        if fmt == "csv":
            return csv_rows(read_all(result))
        with result.stream as f:
            wb = openpyxl.load_workbook(f, read_only=True)
            return [["" if cell is None else cell for cell in row] for row in wb["Maile"].iter_rows(values_only=True)]

    def test_rows_follow_the_order_of_the_ids_in_both_formats(self):
        ids = ("eventy-003", "driverzy-002", "eventy-001")
        for fmt in ("csv", "xlsx"):
            result = self.build(fmt, ids=ids, columns=("company",))
            self.assertEqual(self.table(result, fmt), [["Firma"], ["Gamma"], ["Beta"], ["Alfa"]], fmt)
            self.assertEqual(result.rows, 3, fmt)

    def test_chosen_columns_are_written_in_the_canonical_order(self):
        rows = self.table(self.build(columns=("email", "company", "city")))
        self.assertEqual(rows[0], ["Firma", "Miasto", "E-mail"])  # kolejność pliku, nie kolejność zaznaczania
        self.assertEqual(rows[1], ["Alfa", "London", "firma@example.com"])

    def test_country_name_is_derived_even_when_the_code_is_not_exported(self):
        rows = self.table(self.build(ids=("eventy-001", "driverzy-002"), columns=("company", "country_name")))
        self.assertEqual(rows, [["Firma", "Kraj"], ["Alfa", "Wielka Brytania"], ["Beta", "Polska"]])

    def test_bodies_flag_is_ignored_when_columns_are_given(self):
        self.assertEqual(self.table(self.build(columns=("company",), bodies=True))[0], ["Firma"])
        self.assertIn("Treść maila (tekst)", self.table(self.build(bodies=True))[0])

    def test_unknown_and_empty_column_lists_are_rejected(self):
        with self.assertRaisesRegex(exporter.ExportError, "nope"):
            self.build(columns=("company", "nope"))
        with self.assertRaises(exporter.ExportError):
            self.build(columns=())

    def reported_rows(self, fmt="csv", ids=None):
        result = self.build(fmt, ids=ids)
        result.stream.close()
        return result.rows

    def test_row_count_is_reported_for_every_kind_of_selection(self):
        self.assertEqual(self.reported_rows(), 3)
        self.assertEqual(self.reported_rows("xlsx"), 3)
        self.assertEqual(self.reported_rows(ids=("eventy-001", "nie-ma", "eventy-001")), 1)
        self.assertEqual(self.reported_rows(ids=()), 0)

    def test_catalog_describes_every_exportable_column(self):
        self.assertEqual(set(exporter.COLUMN_INFO), set(exporter.ALL_EXPORT_COLUMNS))
        for label, group in exporter.COLUMN_INFO.values():
            self.assertTrue(label)
            self.assertIn(group, exporter.COLUMN_GROUPS)
        presets = {key: columns for key, _, columns in exporter.COLUMN_PRESETS}
        self.assertEqual(presets["standard"], exporter.EXPORT_COLUMNS)
        self.assertEqual(presets["all"], exporter.ALL_EXPORT_COLUMNS)
        for key, columns in presets.items():
            self.assertTrue(set(columns) <= set(exporter.ALL_EXPORT_COLUMNS), key)
            self.assertTrue(columns, key)


class SelectionEndpointTests(unittest.TestCase):
    def setUp(self):
        self.database = make_db()
        put(self.database,
            sample_row(1, country="GB", company="Alfa", status="wyslano", sent_at="2026-09-29T14:23:11"),
            sample_row(2, segment="driverzy", country="PL", company="Beta"),
            sample_row(3, company="Gamma", email="", body_html="", body_text="", status="pominieto_brak_email"))
        self.app = App(self.database, PAGES, index_html=None, allowed_hosts={f"localhost:{PORT}"})

    def get(self, path, **query):
        return self.app.dispatch(Request("GET", path, query={k: v if isinstance(v, list) else [v] for k, v in query.items()},
                                         headers={"host": f"localhost:{PORT}"}))

    def post(self, payload, content_type="application/json"):
        return self.app.dispatch(Request("POST", "/api/export", body=json.dumps(payload).encode("utf-8"),
                                         headers={"host": f"localhost:{PORT}", "content-type": content_type}))

    def rows_of(self, resp):
        self.assertEqual(resp.status, 200)
        with resp.stream as f:
            return csv_rows(f.read())

    def test_post_returns_the_listed_rows_in_the_listed_order(self):
        resp = self.post({"format": "csv", "ids": ["driverzy-002", "eventy-001"], "columns": ["company", "email"]})
        self.assertEqual(resp.headers["X-Export-Rows"], "2")
        self.assertRegex(resp.headers["Content-Disposition"], r'filename="Maile_\d{4}-\d{2}-\d{2}_\d{4}\.csv"')
        self.assertEqual(self.rows_of(resp), [["Firma", "E-mail"], ["Beta", "firma@example.com"], ["Alfa", "firma@example.com"]])

    def test_post_with_comma_delimiter_and_default_columns(self):
        resp = self.post({"format": "csv", "delimiter": "comma", "ids": ["eventy-001"]})
        with resp.stream as f:
            rows = csv_rows(f.read(), delimiter=",")
        self.assertEqual(tuple(rows[0]), exporter.headers(exporter.EXPORT_COLUMNS))
        self.assertEqual(rows[1][0], "Alfa")
        self.assertEqual(len(rows), 2)

    def test_post_bodies_flag_adds_the_message_columns(self):
        resp = self.post({"format": "csv", "ids": ["eventy-001"], "bodies": True})
        header = self.rows_of(resp)[0]
        self.assertEqual(header[:2], ["Firma", "Segment"])
        self.assertIn("Treść maila (HTML)", header)
        self.assertNotIn("Identyfikator", header)

    def test_post_xlsx(self):
        resp = self.post({"format": "xlsx", "ids": ["eventy-003", "eventy-001"], "columns": ["company"]})
        self.assertIn("spreadsheetml", resp.content_type)
        self.assertEqual(resp.headers["X-Export-Rows"], "2")
        with resp.stream as f:
            wb = openpyxl.load_workbook(f, read_only=True)
            self.assertEqual([r[0] for r in wb["Maile"].iter_rows(values_only=True)], ["Firma", "Gamma", "Alfa"])

    def test_post_validation(self):
        base = {"format": "csv", "ids": ["eventy-001"]}
        too_many = {**base, "ids": [f"x-{n}" for n in range(20_001)]}
        cases = [
            ({"ids": ["eventy-001"]}, "format"),
            ({**base, "format": "pdf"}, "format"),
            ({**base, "format": ["csv"]}, "format"),
            ({**base, "delimiter": "tab"}, "delimiter"),
            ({"format": "csv"}, "ids"),
            ({**base, "ids": "eventy-001"}, "ids"),
            ({**base, "ids": ["eventy-001", 7]}, "ids"),
            ({**base, "ids": [""]}, "ids"),
            ({**base, "ids": []}, "przynajmniej jeden wiersz"),
            (too_many, "Za dużo"),
            ({**base, "columns": "company"}, "columns"),
            ({**base, "columns": ["company", 3]}, "columns"),
            ({**base, "columns": ["company", "nope"]}, "nope"),
            ({**base, "columns": []}, "przynajmniej jedną kolumnę"),
        ]
        for payload, expected in cases:
            resp = self.post(payload)
            self.assertEqual(resp.status, 400, payload if len(str(payload)) < 200 else expected)
            self.assertIn(expected, json.loads(resp.body)["error"])

    def test_post_requires_a_json_content_type(self):
        self.assertEqual(self.post({"format": "csv", "ids": ["eventy-001"]}, content_type="text/plain").status, 415)

    def test_get_understands_the_new_filters(self):
        def count(**query):
            return json.loads(self.get("/api/export/count", **query).body)["count"]

        self.assertEqual(count(country="GB"), 1)
        self.assertEqual(count(country=["GB", "PL"]), 2)
        self.assertEqual(count(country="_none"), 1)
        self.assertEqual(count(has_email="0"), 1)
        self.assertEqual(count(has_body="1"), 2)
        self.assertEqual(count(sent_from="2026-09-29", sent_to="2026-09-29"), 1)
        self.assertEqual(count(sent_from="2026-09-30"), 0)
        self.assertEqual(count(country="PL", has_email="1", segment="driverzy"), 1)
        resp = self.get("/api/export", format="csv", country="GB")
        self.assertEqual(resp.headers["X-Export-Rows"], "1")
        resp.stream.close()

    def test_get_rejects_bad_values_of_the_new_filters(self):
        for query in ({"country": "XX"}, {"has_email": "2"}, {"has_body": "tak"},
                      {"sent_from": "wczoraj"}, {"sent_to": "2026-13-40"}):
            self.assertEqual(self.get("/api/export/count", **query).status, 400, query)

    def test_meta_lists_columns_groups_and_presets(self):
        meta = json.loads(self.get("/api/export/meta").body)
        self.assertEqual([c["name"] for c in meta["columns"]], list(exporter.ALL_EXPORT_COLUMNS))
        self.assertEqual(meta["column_groups"], list(exporter.COLUMN_GROUPS))
        self.assertEqual([p["key"] for p in meta["column_presets"]], ["standard", "contacts", "extended", "minimal", "all"])
        company = next(c for c in meta["columns"] if c["name"] == "company")
        self.assertEqual((company["label"], company["group"]), ("Firma", "Firma i kontakt"))


if __name__ == "__main__":
    unittest.main()
