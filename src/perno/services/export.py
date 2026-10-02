"""Eksport bazy maili do CSV (Excel-PL) lub XLSX. Plik powstaje w całości
w pliku tymczasowym (dane strumieniowane z bazy), dopiero potem idzie do klienta.

Plik ma być czytelny dla człowieka, nie kopią tabeli: nagłówki po polsku („Firma”, „Data wysłania”), status jako
etykieta („Wysłano”), kraj jako nazwa, daty bez sekund, a domyślnie tylko kolumny, które się czyta (reszta — pola
techniczne i treść maili — jest do dobrania na stronie Eksport)."""

from __future__ import annotations

import csv
import tempfile
from dataclasses import dataclass
from datetime import datetime
from typing import BinaryIO, Optional

from perno import db as dbmod
from perno.db import RowFilter
from perno.domain import countries
from perno.domain import statuses as st

BODY_COLUMNS = ("body_html", "body_text")
# `country_name` nie jest kolumną bazy — to polska nazwa kraju wyliczana z kodu, żeby plik był czytelny w Excelu.
DERIVED_COLUMNS = ("country_name",)

# Wszystkie kolumny, jakie można wyeksportować, w stałej kolejności pliku: najpierw to, co czyta człowiek,
# na końcu pola techniczne. Kolejność, w jakiej ktoś zaznaczył kolumny, nie ma znaczenia.
ALL_EXPORT_COLUMNS = (
    "company", "segment", "country_name", "country", "city", "address", "email", "contact_person", "website",
    "status", "sent_at", "generated_at", "subject", "description", "error",
    "body_html", "body_text",
    "id", "row_number", "tag", "mode", "attachments", "inline_images", "retry_count", "updated_at",
)
# Zestaw domyślny: tyle, żeby plik dało się przeczytać bez przewijania na boki.
STANDARD_COLUMNS = ("company", "segment", "country_name", "city", "email", "website", "status", "sent_at", "subject")
EXPORT_COLUMNS = STANDARD_COLUMNS

# Dla nagłówków pliku i okna wyboru kolumn: nazwa → (nagłówek po polsku, grupa).
COLUMN_GROUPS = ("Firma i kontakt", "Lokalizacja", "Wysyłka", "Treść maili", "Techniczne")
COLUMN_INFO = {
    "company": ("Firma", "Firma i kontakt"),
    "segment": ("Segment", "Firma i kontakt"),
    "contact_person": ("Osoba kontaktowa", "Firma i kontakt"),
    "email": ("E-mail", "Firma i kontakt"),
    "website": ("Strona www", "Firma i kontakt"),
    "description": ("Opis działalności", "Firma i kontakt"),
    "country_name": ("Kraj", "Lokalizacja"),
    "country": ("Kod kraju", "Lokalizacja"),
    "city": ("Miasto", "Lokalizacja"),
    "address": ("Adres", "Lokalizacja"),
    "status": ("Status", "Wysyłka"),
    "sent_at": ("Data wysłania", "Wysyłka"),
    "generated_at": ("Data wygenerowania", "Wysyłka"),
    "subject": ("Temat maila", "Wysyłka"),
    "error": ("Błąd wysyłki", "Wysyłka"),
    "body_html": ("Treść maila (HTML)", "Treść maili"),
    "body_text": ("Treść maila (tekst)", "Treść maili"),
    "id": ("Identyfikator", "Techniczne"),
    "row_number": ("Numer w segmencie", "Techniczne"),
    "tag": ("Tag", "Techniczne"),
    "mode": ("Forma zwrotu", "Techniczne"),
    "attachments": ("Załączniki", "Techniczne"),
    "inline_images": ("Obrazy w treści", "Techniczne"),
    "retry_count": ("Liczba prób wysyłki", "Techniczne"),
    "updated_at": ("Data zmiany", "Techniczne"),
}
# Gotowe zestawy kolumn: (klucz, etykieta, kolumny).
COLUMN_PRESETS = (
    ("standard", "Standardowy", STANDARD_COLUMNS),
    ("contacts", "Kontakty", ("company", "segment", "country_name", "city", "email", "website")),
    ("extended", "Rozszerzony", ("company", "segment", "country_name", "city", "address", "email", "contact_person",
                                 "website", "status", "sent_at", "generated_at", "subject", "description", "error")),
    ("minimal", "Minimalny", ("company", "email")),
    ("all", "Wszystko (z treścią maili i polami technicznymi)", ALL_EXPORT_COLUMNS),
)

XLSX_MAX_ROWS = 1_048_575
XLSX_MAX_CELL = 32_767
XLSX_MAX_URL = 2000  # Excel odrzuca dłuższe odnośniki
DATETIME_TEXT = "%Y-%m-%d %H:%M"
DATETIME_XLSX = "yyyy-mm-dd hh:mm"
CSV_DANGEROUS = ("=", "+", "-", "@", "\t", "\r")

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
DELIMITERS = {"semicolon": ";", "comma": ","}
_WIDTHS = {
    "company": 32, "segment": 11, "country_name": 18, "country": 10, "city": 18, "address": 34, "email": 32,
    "contact_person": 22, "website": 36, "status": 16, "sent_at": 17, "generated_at": 19, "subject": 62,
    "description": 56, "error": 40, "body_html": 60, "body_text": 60, "id": 16, "row_number": 11, "tag": 18,
    "mode": 13, "attachments": 18, "inline_images": 20, "retry_count": 11, "updated_at": 17,
}


class ExportError(Exception):
    pass


@dataclass(frozen=True)
class ExportSpec:
    fmt: str
    flt: RowFilter = RowFilter()
    bodies: bool = False
    delimiter: str = ";"
    columns: Optional[tuple] = None  # None = zestaw standardowy (z treścią maili, gdy `bodies`); podane `bodies` ignoruje


@dataclass
class ExportResult:
    stream: BinaryIO
    length: int
    filename: str
    content_type: str
    rows: int = 0


def resolve_columns(spec: ExportSpec) -> tuple:
    """Kolumny pliku: zestaw standardowy albo wybrane przez użytkownika (zawsze w kolejności kanonicznej)."""
    if spec.columns is None:
        return tuple(c for c in ALL_EXPORT_COLUMNS if c in STANDARD_COLUMNS or (spec.bodies and c in BODY_COLUMNS))
    unknown = [name for name in spec.columns if name not in ALL_EXPORT_COLUMNS]
    if unknown:
        raise ExportError(f"Nieznana kolumna: {unknown[0]}.")
    chosen = tuple(name for name in ALL_EXPORT_COLUMNS if name in spec.columns)
    if not chosen:
        raise ExportError("Wybierz przynajmniej jedną kolumnę.")
    return chosen


def headers(columns: tuple) -> tuple:
    """Nagłówki pliku (po polsku) dla podanych kolumn."""
    return tuple(COLUMN_INFO[name][0] for name in columns)


def status_label(code) -> str:
    return st.STATUS_META.get(code, (code or "",))[0]


class _Utf8Sink:
    def __init__(self, fileobj: BinaryIO):
        self._file = fileobj

    def write(self, text: str) -> int:
        self._file.write(text.encode("utf-8"))
        return len(text)


def _csv_cell(name: str, value):
    if value is None:
        return ""
    if name == "status":
        value = status_label(value)
    if isinstance(value, datetime):
        return value.strftime(DATETIME_TEXT)
    if isinstance(value, str) and value.startswith(CSV_DANGEROUS):
        return "'" + value
    return value


def _rows(db: dbmod.Database, spec: ExportSpec, columns: tuple):
    """Wiersze z bazy (tylko prawdziwe kolumny) uzupełnione o kolumny wyliczane."""
    real = tuple(c for c in columns if c not in DERIVED_COLUMNS)
    if "country_name" in columns and "country" not in real:
        real += ("country",)  # nazwa kraju jest liczona z kodu, więc kod trzeba pobrać, nawet gdy nie trafia do pliku
    for row in db.iter_rows(spec.flt, real):
        if "country_name" in columns:
            row["country_name"] = countries.country_name(row.get("country"))
        yield row


def _write_csv(db: dbmod.Database, spec: ExportSpec, columns: tuple, out: BinaryIO) -> int:
    out.write(b"\xef\xbb\xbf")
    writer = csv.writer(_Utf8Sink(out), delimiter=spec.delimiter, lineterminator="\r\n")
    writer.writerow(headers(columns))
    count = 0
    for row in _rows(db, spec, columns):
        writer.writerow([_csv_cell(name, row[name]) for name in columns])
        count += 1
    return count


class _XlsxStyles:
    """Style komórek tworzone raz (openpyxl łączy identyczne), żeby plik nie puchł."""

    def __init__(self):
        from openpyxl.styles import Alignment, Font, PatternFill

        self.header_font = Font(bold=True, color="FFFFFFFF")
        self.header_fill = PatternFill("solid", start_color="FF1F3A5F")
        self.header_align = Alignment(vertical="center", wrap_text=True)
        self.body_align = Alignment(vertical="top", wrap_text=True)  # długi tekst zawija się w komórce, całość od góry
        self.link_font = Font(color="FF1155CC", underline="single")
        self._status = {}

    def status_font(self, code):
        from openpyxl.styles import Font

        if code not in self._status:
            color = st.STATUS_META.get(code, ("", "#6b7280"))[1].lstrip("#")
            self._status[code] = Font(bold=True, color="FF" + color.upper())
        return self._status[code]


def _xlsx_cell(ws, styles: _XlsxStyles, name: str, value):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if value is None:
        return None
    cell = None
    if name == "status":
        cell = WriteOnlyCell(ws, value=status_label(value))
        cell.font = styles.status_font(value)
    elif isinstance(value, datetime):
        cell = WriteOnlyCell(ws, value=value)
        cell.number_format = DATETIME_XLSX
    elif isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:XLSX_MAX_CELL]
        cell = WriteOnlyCell(ws, value=value)
        if name == "website" and value.lower().startswith(("http://", "https://")) and len(value) <= XLSX_MAX_URL:
            cell.hyperlink = value
            cell.font = styles.link_font
        elif value.startswith("="):
            cell.data_type = "s"  # inaczej openpyxl zapisałby to jako formułę
    else:
        cell = WriteOnlyCell(ws, value=value)
    cell.alignment = styles.body_align
    return cell


def _write_xlsx(db: dbmod.Database, spec: ExportSpec, columns: tuple, out: BinaryIO) -> int:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    if db.count_rows(spec.flt) > XLSX_MAX_ROWS:
        raise ExportError("Za dużo wierszy na jeden arkusz XLSX — wybierz format CSV.")

    styles = _XlsxStyles()
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Maile")
    for index, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = _WIDTHS.get(name, 18)
    ws.freeze_panes = "B2" if columns[0] == "company" else "A2"  # nagłówek i nazwa firmy zostają na ekranie
    ws.row_dimensions[1].height = 26
    # wydruk: poziomo, wszystkie kolumny na szerokość strony, nagłówek na każdej stronie
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"

    header = []
    for label in headers(columns):
        cell = WriteOnlyCell(ws, value=label)
        cell.font, cell.fill, cell.alignment = styles.header_font, styles.header_fill, styles.header_align
        header.append(cell)
    ws.append(header)

    count = 0
    for row in _rows(db, spec, columns):
        ws.append([_xlsx_cell(ws, styles, name, row[name]) for name in columns])
        count += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count + 1}"
    wb.save(out)
    return count


def build_export(db: dbmod.Database, spec: ExportSpec, now: Optional[datetime] = None) -> ExportResult:
    if spec.fmt not in CONTENT_TYPES:
        raise ExportError(f"Nieznany format: {spec.fmt!r}.")
    columns = resolve_columns(spec)
    out = tempfile.TemporaryFile()  # noqa: SIM115 — plik zamykamy po wysłaniu odpowiedzi albo przy błędzie (niżej)
    try:
        rows = (_write_csv if spec.fmt == "csv" else _write_xlsx)(db, spec, columns, out)
        length = out.tell()
        out.seek(0)
    except BaseException:
        out.close()
        raise
    stamp = (now or datetime.now()).strftime("%Y-%m-%d_%H%M")
    return ExportResult(out, length, f"Maile_{stamp}.{spec.fmt}", CONTENT_TYPES[spec.fmt], rows)
